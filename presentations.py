# -*- coding: utf-8 -*-
"""Babata: PowerPoint prezentatsiya yaratish (python-pptx)."""

from __future__ import annotations

import os
from typing import List, Optional, Tuple

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    _HAVE_PPTX = True
except ImportError:
    _HAVE_PPTX = False


def create_presentation(
    topic: str,
    slides_content: Optional[List[dict]] = None,
    output_path: Optional[str] = None,
    lang: str = "en",
) -> Tuple[bool, str]:
    """
    PowerPoint (.pptx) yaratish.
    slides_content: [{"title": "...", "content": "..."}, ...]
    Bo'sh bo'lsa — faqat title slayd yaratiladi.
    Qaytadi: (success, message yoki path).
    """
    if not _HAVE_PPTX:
        return False, "python-pptx not installed. Run: pip install python-pptx"

    topic = (topic or "Presentation").strip()[:200]
    if not output_path:
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        if not os.path.isdir(desktop):
            desktop = os.path.expanduser("~")
        safe_name = "".join(c for c in topic.replace(" ", "_")[:50] if c.isalnum() or c in "_-") or "NeuroKey_Presentation"
        output_path = os.path.join(desktop, f"{safe_name}.pptx")

    try:
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        if slides_content and len(slides_content) > 0:
            for s in slides_content:
                title = (s.get("title") or "").strip()[:200] or "Slide"
                content = (s.get("content") or "").strip()[:1500] or ""
                slide_layout = prs.slide_layouts[6]  # Blank
                slide = prs.slides.add_slide(slide_layout)
                left, top, width, height = Inches(0.5), Inches(0.5), Inches(9), Inches(1)
                title_box = slide.shapes.add_textbox(left, top, width, height)
                tf = title_box.text_frame
                p = tf.paragraphs[0]
                p.text = title
                p.font.size = Pt(28)
                p.font.bold = True
                if content:
                    body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
                    tf2 = body_box.text_frame
                    tf2.word_wrap = True
                    p2 = tf2.paragraphs[0]
                    p2.text = content[:800]
                    p2.font.size = Pt(18)
        else:
            slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(slide_layout)
            box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
            tf = box.text_frame
            tf.paragraphs[0].text = topic
            tf.paragraphs[0].font.size = Pt(36)
            tf.paragraphs[0].font.bold = True

        prs.save(output_path)
        return True, output_path
    except Exception as e:
        return False, str(e)


def create_presentation_from_outline(outline: str, topic: str = "", output_path: Optional[str] = None) -> Tuple[bool, str]:
    """
    outline: "1. Intro\n2. Main points\n3. Conclusion" yoki "Slide 1: ...\nSlide 2: ..."
    Har bir qatordan slayd yaratadi.
    """
    lines = [l.strip() for l in (outline or "").splitlines() if l.strip()]
    slides = []
    for line in lines:
        if ":" in line:
            parts = line.split(":", 1)
            slides.append({"title": parts[0].strip()[:100], "content": (parts[1] or "").strip()[:800]})
        else:
            slides.append({"title": line[:100], "content": ""})
    if not slides and topic:
        slides = [{"title": topic, "content": ""}]
    return create_presentation(topic or "Presentation", slides_content=slides, output_path=output_path, lang="en")
